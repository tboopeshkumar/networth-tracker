"""Build a LOCAL-ONLY demo fixture from an .xlsx export of the tracker.

The fixture mimics what the Sheets API returns (values:batchGet with
UNFORMATTED_VALUE / SERIAL_NUMBER, plus a FORMULA render) so the app can be
exercised on localhost without Google credentials.

    python3 tools/make_fixture.py ~/Downloads/NetworthTracker.xlsx

Writes demo/fixture.json, which is gitignored and blocked by the pre-commit
hook. Never commit it: it contains your real data.
"""
import datetime
import json
import pathlib
import sys

import openpyxl

EPOCH = datetime.datetime(1899, 12, 30)


def cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return (v - EPOCH).total_seconds() / 86400
    if isinstance(v, datetime.date):
        return (datetime.datetime(v.year, v.month, v.day) - EPOCH).days
    if hasattr(v, "text"):  # openpyxl ArrayFormula
        return v.text
    if not isinstance(v, (str, int, float, bool)):
        return str(v)
    return v


def grid(ws):
    rows = []
    for r in ws.iter_rows(values_only=True):
        row = [cell(v) for v in r]
        while row and row[-1] == "":
            row.pop()
        rows.append(row)
    while rows and not rows[-1]:
        rows.pop()
    return rows


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    src = pathlib.Path(sys.argv[1]).expanduser()
    values_wb = openpyxl.load_workbook(src, data_only=True)
    formula_wb = openpyxl.load_workbook(src, data_only=False)

    fixture = {"sheets": [], "values": {}, "formulas": {}}
    for i, name in enumerate(values_wb.sheetnames):
        fixture["sheets"].append({"title": name, "sheetId": 1000 + i})
        fixture["values"][name] = grid(values_wb[name])
        fixture["formulas"][name] = grid(formula_wb[name])

    out = pathlib.Path(__file__).resolve().parent.parent / "demo" / "fixture.json"
    out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps(fixture, ensure_ascii=False))
    print(f"wrote {out} ({len(fixture['sheets'])} tabs) — local only, do not commit")


if __name__ == "__main__":
    main()
